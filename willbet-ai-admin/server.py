from http.server import ThreadingHTTPServer,SimpleHTTPRequestHandler
from pathlib import Path
import io,json,os,argparse
from openpyxl import Workbook,load_workbook
os.chdir(Path(__file__).parent)
HEAD=['知识名称','回复正文']
class Handler(SimpleHTTPRequestHandler):
 def end_headers(self):
  self.send_header('Cache-Control','no-store, no-cache, must-revalidate')
  super().end_headers()
 def send_head(self):
  # SimpleHTTPRequestHandler otherwise returns 304 for cached static files.
  if 'If-Modified-Since' in self.headers:del self.headers['If-Modified-Since']
  return super().send_head()
 def do_GET(self):
  if self.path=='/api/template':
   w=Workbook();s=w.active;s.title='知识导入';s.append(HEAD)
   for col in s.columns:s.column_dimensions[col[0].column_letter].width=32
   b=io.BytesIO();w.save(b);self.send_response(200);self.send_header('Content-Type','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet');self.send_header('Content-Disposition','attachment; filename=knowledge-template.xlsx');self.end_headers();self.wfile.write(b.getvalue())
  else:super().do_GET()
 def do_POST(self):
  if self.path!='/api/parse':self.send_error(404);return
  try:
   length=int(self.headers.get('Content-Length',0))
   if length>10_000_000:raise ValueError('文件不能超过 10 MB')
   w=load_workbook(io.BytesIO(self.rfile.read(length)),read_only=True,data_only=True);rows=list(w.active.values)
   if not rows or list(rows[0])!=HEAD:raise ValueError('表头与模板不一致，请下载模板后重新上传')
   result=[dict(zip(['name','body'],[str(v).strip() if v is not None else '' for v in r[:2]])) for r in rows[1:] if any(v is not None for v in r)]
   self.send_response(200);payload={'rows':result}
  except Exception as e:self.send_response(400);payload={'error':str(e)}
  self.send_header('Content-Type','application/json');self.end_headers();self.wfile.write(json.dumps(payload,ensure_ascii=False).encode())
parser=argparse.ArgumentParser(description='AIBot Admin 本地预览服务')
parser.add_argument('--port',type=int,default=4180)
args=parser.parse_args()
server=ThreadingHTTPServer(('127.0.0.1',args.port),Handler)
print(f'Local: http://127.0.0.1:{args.port}',flush=True)
server.serve_forever()
